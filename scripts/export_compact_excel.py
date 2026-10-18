from __future__ import annotations

import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


PROJECT_ROOT = Path(__file__).resolve().parents[1]
RUN_DIR = PROJECT_ROOT / "data" / "annotations" / "gpt6_luna_compact_10000"
OUTPUT_PATH = PROJECT_ROOT / "outputs" / "gpt6_luna_compact_10000_review.xlsx"
VALIDATION_PATH = PROJECT_ROOT / "outputs" / "gpt6_luna_compact_10000_review.validation.json"

FONT = "Arial"
DARK_BLUE = "1F4E78"
MID_BLUE = "5B9BD5"
LIGHT_BLUE = "DDEBF7"
LIGHT_GRAY = "F2F2F2"
LIGHT_RED = "FCE4D6"
LIGHT_GREEN = "E2F0D9"
LIGHT_AMBER = "FFF2CC"
WHITE = "FFFFFF"
TEXT = "1F1F1F"
THIN_GRAY = Side(style="thin", color="D9E1F2")


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [
        json.loads(line)
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]


def flatten(row: dict[str, Any]) -> list[Any]:
    annotation = row["annotation"]
    return [
        row["article_id"],
        row.get("published_at", ""),
        row.get("month", ""),
        row.get("article_type", ""),
        row.get("section", ""),
        row.get("headline", ""),
        row.get("abstract", ""),
        annotation.get("relevance"),
        annotation.get("primary_topic"),
        annotation.get("tone"),
        annotation.get("topic_evidence", ""),
        annotation.get("tone_evidence", ""),
        annotation.get("needs_review"),
        ", ".join(annotation.get("review_flags") or []),
        ", ".join(row.get("validation_issues") or []),
        ", ".join(row.get("normalization_actions") or []),
        row.get("duplicate_group_id", ""),
    ]


HEADERS = [
    "article_id",
    "published_at",
    "month",
    "article_type",
    "section",
    "headline",
    "abstract",
    "relevance",
    "primary_topic",
    "tone",
    "topic_evidence",
    "tone_evidence",
    "needs_review",
    "review_flags",
    "validation_issues",
    "normalization_actions",
    "duplicate_group_id",
]


def style_title(ws, title: str) -> None:
    ws.sheet_view.showGridLines = False
    ws["A2"] = title
    ws["A2"].font = Font(name=FONT, size=15, bold=True, color=TEXT)
    ws["A3"].border = Border(bottom=Side(style="thin", color=DARK_BLUE))


def style_header(ws, row: int, start_col: int, end_col: int) -> None:
    for cell in ws.iter_cols(min_col=start_col, max_col=end_col, min_row=row, max_row=row):
        item = cell[0]
        item.fill = PatternFill("solid", fgColor=DARK_BLUE)
        item.font = Font(name=FONT, size=10, bold=True, color=WHITE)
        item.alignment = Alignment(horizontal="center", vertical="center")
        item.border = Border(bottom=THIN_GRAY)


def add_data_sheet(wb: Workbook, name: str, rows: Iterable[dict[str, Any]], table_name: str) -> int:
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.append(HEADERS)
    count = 0
    for row in rows:
        ws.append(flatten(row))
        count += 1
    style_header(ws, 1, 1, len(HEADERS))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = Font(name=FONT, size=9, color=TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    widths = {
        "A": 38,
        "B": 24,
        "C": 11,
        "D": 12,
        "E": 18,
        "F": 48,
        "G": 80,
        "H": 12,
        "I": 24,
        "J": 8,
        "K": 42,
        "L": 42,
        "M": 13,
        "N": 28,
        "O": 34,
        "P": 38,
        "Q": 25,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:Q{ws.max_row}"
    if count:
        table = Table(displayName=table_name, ref=f"A1:Q{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)
        ws.conditional_formatting.add(
            f"J2:J{ws.max_row}",
            CellIsRule(operator="equal", formula=["1"], fill=PatternFill("solid", fgColor=LIGHT_GREEN)),
        )
        ws.conditional_formatting.add(
            f"J2:J{ws.max_row}",
            CellIsRule(operator="equal", formula=["-1"], fill=PatternFill("solid", fgColor=LIGHT_RED)),
        )
        ws.conditional_formatting.add(
            f"O2:O{ws.max_row}",
            FormulaRule(formula=["LEN($O2)>0"], fill=PatternFill("solid", fgColor=LIGHT_AMBER)),
        )
    return count


def add_summary_sheet(
    wb: Workbook,
    summary: dict[str, Any],
    pilot: dict[str, Any],
) -> None:
    ws = wb.active
    ws.title = "Summary"
    style_title(ws, "GPT economic news annotation summary")
    ws["A4"] = "Run completed 2026-10-15. Source pool was prefiltered as relevant; GPT relevance is a second-stage content judgment."
    ws["A4"].font = Font(name=FONT, size=10, italic=True, color="666666")

    metrics = [
        ("Selected articles", summary["selected_total"]),
        ("GPT relevant", summary["relevance"]["relevant"]),
        ("GPT irrelevant", summary["relevance"]["irrelevant"]),
        ("GPT uncertain", summary["relevance"]["uncertain"]),
        ("Review queue", summary["review_queue_count"]),
        ("Estimated API cost (USD)", summary["cost_estimate_usd"]["total_run_including_repeats_and_retries_approx"]),
    ]
    ws.append([])
    ws.append(["Metric", "Value"])
    for label, value in metrics:
        ws.append([label, value])
    style_header(ws, 6, 1, 2)
    ws["B12"].number_format = '$0.00'

    ws["D6"] = "Relevance"
    ws["E6"] = "Count"
    ws["F6"] = "Share"
    style_header(ws, 6, 4, 6)
    for index, label in enumerate(["relevant", "irrelevant", "uncertain"], 7):
        value = summary["relevance"][label]
        ws.cell(index, 4, label)
        ws.cell(index, 5, value)
        ws.cell(index, 6, value / summary["selected_total"])
        ws.cell(index, 6).number_format = "0.0%"

    relevant_total = summary["relevance"]["relevant"]
    ws["A15"] = "Tone among GPT-relevant articles"
    ws["A16"] = "Tone"
    ws["B16"] = "Count"
    ws["C16"] = "Share"
    style_header(ws, 16, 1, 3)
    for index, label in enumerate(["-1", "0", "1"], 17):
        value = summary["tone_relevant_only"][label]
        ws.cell(index, 1, int(label))
        ws.cell(index, 2, value)
        ws.cell(index, 3, value / relevant_total)
        ws.cell(index, 3).number_format = "0.0%"

    ws["E15"] = "Primary topic among GPT-relevant articles"
    ws["E16"] = "Primary topic"
    ws["F16"] = "Count"
    ws["G16"] = "Share"
    style_header(ws, 16, 5, 7)
    topic_order = [
        "cost_of_living_income",
        "labor_market",
        "household_economy",
        "macro_growth",
        "economic_policy",
    ]
    for index, label in enumerate(topic_order, 17):
        value = summary["primary_topic_relevant_only"][label]
        ws.cell(index, 5, label)
        ws.cell(index, 6, value)
        ws.cell(index, 7, value / relevant_total)
        ws.cell(index, 7).number_format = "0.0%"

    ws["A24"] = "Quality checks"
    ws["A25"] = "Check"
    ws["B25"] = "Result"
    style_header(ws, 25, 1, 2)
    evidence_issues = (
        summary["validation_issues"].get("topic_evidence_not_exact_substring", 0)
        + summary["validation_issues"].get("tone_evidence_not_exact_substring", 0)
    )
    checks = [
        ("Successful labels", summary["combined_successful"]),
        ("Missing labels", summary["combined_missing"]),
        ("API failures after retry", summary["production_failed"]),
        ("article_id values corrected from manifest", summary["normalization"]["article_id_corrected_from_request_manifest"]),
        ("Leading JSON accepted with trailing response text", summary["validation_issues"].get("response_trailing_text", 0)),
        ("Evidence quotes needing manual review", evidence_issues),
        ("Blind-repeat pairs", pilot["blind_repeat_agreement"]["paired"]),
        ("Full label tuple repeat agreement", pilot["blind_repeat_agreement"]["full_label_tuple"]["rate"]),
        ("Topic repeat agreement when both relevant", pilot["blind_repeat_agreement"]["primary_topic_when_both_relevant"]["rate"]),
        ("Tone repeat agreement when both relevant", pilot["blind_repeat_agreement"]["tone_when_both_relevant"]["rate"]),
    ]
    for index, (label, value) in enumerate(checks, 26):
        ws.cell(index, 1, label)
        ws.cell(index, 2, value)
        if "agreement" in label:
            ws.cell(index, 2).number_format = "0.0%"
    ws["D25"] = "Method note"
    ws["E25"] = "Detail"
    style_header(ws, 25, 4, 5)
    notes = [
        ("Tone", "+1 favorable, -1 adverse, 0 neutral/mixed/insufficient."),
        ("Topics", "Exactly one primary topic for GPT-relevant articles; five consolidated classes."),
        ("Sampling", "10,000 new IDs; old 5,000 IDs and their duplicate groups excluded; one record per new duplicate group."),
        ("Time balance", "Most months contain 49-50 articles. 2026-01 and 2026-10 were capacity constrained at 23 and 24."),
        ("Accuracy", "Blind repeats measure consistency, not accuracy against an independent human gold standard."),
        ("Evidence", "Evidence text was not auto-rewritten; non-exact quotes remain in the review queue."),
    ]
    for index, (label, detail) in enumerate(notes, 26):
        ws.cell(index, 4, label)
        ws.cell(index, 5, detail)

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=7):
        for cell in row:
            if not cell.font.name:
                cell.font = Font(name=FONT, size=10, color=TEXT)
            cell.alignment = Alignment(vertical="center", wrap_text=False)
    ws.column_dimensions["A"].width = 43
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 23
    ws.column_dimensions["E"].width = 88
    ws.column_dimensions["F"].width = 14
    ws.column_dimensions["G"].width = 14
    ws.freeze_panes = "A6"
    ws.sheet_properties.tabColor = DARK_BLUE


def add_monthly_sheet(wb: Workbook, rows: list[dict[str, Any]]) -> None:
    monthly: dict[str, Counter] = defaultdict(Counter)
    for row in rows:
        annotation = row["annotation"]
        month = row["month"]
        monthly[month]["total"] += 1
        relevance = annotation["relevance"]
        monthly[month][relevance] += 1
        if relevance == "relevant":
            monthly[month][f"tone_{annotation['tone']}"] += 1
    ws = wb.create_sheet("Monthly counts")
    ws.sheet_view.showGridLines = False
    headers = ["month", "total", "relevant", "irrelevant", "uncertain", "tone_-1", "tone_0", "tone_1"]
    ws.append(headers)
    for month in sorted(monthly):
        counts = monthly[month]
        ws.append([month] + [counts[key] for key in headers[1:]])
    style_header(ws, 1, 1, len(headers))
    for column in "ABCDEFGH":
        ws.column_dimensions[column].width = 14
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT, size=9, color=TEXT)
            cell.alignment = Alignment(horizontal="right" if cell.column > 1 else "left")
    ws.freeze_panes = "B2"
    table = Table(displayName="MonthlyCountsTable", ref=f"A1:H{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)


def add_audit_sheet(
    wb: Workbook,
    selection: dict[str, Any],
    summary: dict[str, Any],
    pilot: dict[str, Any],
) -> None:
    ws = wb.create_sheet("Audit")
    style_title(ws, "Selection, validation and cost audit")
    ws["A5"] = "Field"
    ws["B5"] = "Value"
    style_header(ws, 5, 1, 2)
    items = [
        ("Selection seed", selection["selection_seed"]),
        ("Source rows", selection["source_rows"]),
        ("Old IDs excluded", selection["excluded_old_ids"]),
        ("Old duplicate-group members excluded", selection["excluded_old_duplicate_groups"]),
        ("New duplicate-group members removed", selection["removed_new_duplicate_group_members"]),
        ("Selected IDs", selection["selected_unique_ids"]),
        ("Selected duplicate groups", selection["selected_unique_duplicate_groups"]),
        ("Old ID overlap", selection["old_id_overlap"]),
        ("Old duplicate-group overlap", selection["old_duplicate_group_overlap"]),
        ("Pilot unique", selection["pilot_unique_count"]),
        ("Pilot blind repeats", selection["pilot_repeat_count"]),
        ("Pilot full tuple repeat agreement", pilot["blind_repeat_agreement"]["full_label_tuple"]["rate"]),
        ("Prompt input tokens, final responses", summary["usage_all_10000"]["input_tokens"]),
        ("Cached input tokens, final responses", summary["usage_all_10000"]["cached_tokens"]),
        ("Cache write tokens, final responses", summary["usage_all_10000"]["cache_write_tokens"]),
        ("Output tokens, final responses", summary["usage_all_10000"]["output_tokens"]),
        ("Production API calls including retries", summary["estimated_api_calls"]["production_including_validation_retries"]),
        ("Pilot API calls including repeats/retries", summary["estimated_api_calls"]["pilot_including_repeats_and_parse_retries"]),
        ("Estimated total API cost (USD)", summary["cost_estimate_usd"]["total_run_including_repeats_and_retries_approx"]),
        ("Pricing rate date", summary["cost_estimate_usd"]["rate_snapshot_date"]),
        ("Pricing source", summary["cost_estimate_usd"]["official_rate_url"]),
    ]
    for row_index, (label, value) in enumerate(items, 6):
        ws.cell(row_index, 1, label)
        ws.cell(row_index, 2, value)
        if "agreement" in label:
            ws.cell(row_index, 2).number_format = "0.0%"
        if "cost" in label.lower():
            ws.cell(row_index, 2).number_format = '$0.00'
    pricing_cell = ws.cell(6 + len(items) - 1, 2)
    pricing_cell.hyperlink = str(pricing_cell.value)
    pricing_cell.style = "Hyperlink"
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=2):
        for cell in row:
            if cell.style != "Hyperlink":
                cell.font = Font(name=FONT, size=10, color=TEXT, bold=cell.font.bold)
            cell.alignment = Alignment(vertical="center")
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 88
    ws.freeze_panes = "A6"


def main() -> None:
    rows = read_jsonl(RUN_DIR / "annotations_10000.jsonl")
    review = read_jsonl(RUN_DIR / "review_queue.jsonl")
    summary = read_json(RUN_DIR / "production_summary.json")
    pilot = read_json(RUN_DIR / "pilot_summary.json")
    selection = read_json(RUN_DIR / "selection_audit.json")
    relevant = [row for row in rows if row["annotation"]["relevance"] == "relevant"]

    wb = Workbook()
    add_summary_sheet(wb, summary, pilot)
    all_count = add_data_sheet(wb, "All annotations", rows, "AllAnnotationsTable")
    relevant_count = add_data_sheet(wb, "GPT relevant", relevant, "RelevantAnnotationsTable")
    review_count = add_data_sheet(wb, "Review queue", review, "ReviewQueueTable")
    add_monthly_sheet(wb, rows)
    add_audit_sheet(wb, selection, summary, pilot)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)

    check = load_workbook(OUTPUT_PATH, read_only=False, data_only=False)
    expected_sheets = ["Summary", "All annotations", "GPT relevant", "Review queue", "Monthly counts", "Audit"]
    checks = {
        "path": str(OUTPUT_PATH.resolve()),
        "file_size_bytes": OUTPUT_PATH.stat().st_size,
        "sheet_names": check.sheetnames,
        "sheet_names_match": check.sheetnames == expected_sheets,
        "all_annotation_rows": check["All annotations"].max_row - 1,
        "gpt_relevant_rows": check["GPT relevant"].max_row - 1,
        "review_queue_rows": check["Review queue"].max_row - 1,
        "monthly_rows": check["Monthly counts"].max_row - 1,
        "expected_counts_match": (
            all_count == 10000
            and relevant_count == summary["relevance"]["relevant"]
            and review_count == summary["review_queue_count"]
            and check["All annotations"].max_row - 1 == all_count
            and check["GPT relevant"].max_row - 1 == relevant_count
            and check["Review queue"].max_row - 1 == review_count
        ),
        "formula_error_strings_found": 0,
        "visual_render": "Unavailable: artifact-tool runtime/loader and local office renderer were not present; layout was validated structurally with openpyxl.",
    }
    for ws in check.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(
                    token in cell.value
                    for token in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")
                ):
                    checks["formula_error_strings_found"] += 1
    check.close()
    VALIDATION_PATH.write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    if not checks["sheet_names_match"] or not checks["expected_counts_match"] or checks["formula_error_strings_found"]:
        raise RuntimeError(f"Workbook validation failed: {checks}")
    print(json.dumps(checks, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
